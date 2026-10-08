"""Build the manuscript as a Word document in the house style of the reference paper: same styles, page geometry,
header/footer, author block, raised blue citation numerals, blue cross-references, caption-below tables and figures,
full-width diagrams, and an algorithm block. Every number comes from results/*.json or paper/numbers.tex.
    python -m experiments.build_docx
"""
from __future__ import annotations

import copy
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

from experiments.docx_text import ABSTRACT, BACK, BODY, KEYWORDS, REFS, RUNNING_TITLE, TITLE, context
from src.utils import ROOT

REF_DOCX = str(ROOT / "paper" / "_template.docx")   # journal-style template (local, git-ignored); supplies styles, page geometry, author block
OUT = ROOT / "paper" / "When_does_uncertainty_pay_Jerath_Jagadeesan.docx"
TNR, CORBEL, CAMBRIA = "Times New Roman", "Corbel", "Cambria"
MINUS = "\u2212"
BLUE = RGBColor(0x00, 0x00, 0xFF)
TEXT_W = 8005                      # text-column width of the template, in twentieths of a point


# ----------------------------------------------------------------------------------------------- numbers
def load_numbers() -> dict[str, str]:
    nums = {}
    for fn in ("numbers.tex", "numbers2.tex"):
        p = ROOT / "paper" / fn
        if not p.exists():
            continue
        for ln in p.read_text(encoding="utf-8").splitlines():
            if ln.startswith("\\newcommand{\\"):
                nums[ln[len("\\newcommand{\\"):ln.index("}")]] = ln[ln.index("}{") + 2:-1].replace("$-$", MINUS)
    return nums


NUMS = load_numbers()


def subst(text: str) -> str:
    def rep(m):
        if m.group(1) not in NUMS:
            raise KeyError(f"unknown macro {{{{{m.group(1)}}}}}")
        return NUMS[m.group(1)]
    return re.sub(r"\{\{(\w+)\}\}", rep, text)


# ----------------------------------------------------------------------------------------------- citations
class Cites:
    def __init__(self):
        self.order: list[str] = []

    def num(self, key: str) -> int:
        if key not in REFS:
            raise KeyError(f"reference {key} not defined")
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def render(self, keys: list[str]) -> str:
        ns = sorted(self.num(k) for k in keys)
        parts, i = [], 0
        while i < len(ns):
            j = i
            while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                j += 1
            parts.append(str(ns[i]) if j == i else (f"{ns[i]},{ns[j]}" if j == i + 1 else f"{ns[i]}\u2013{ns[j]}"))
            i = j + 1
        return ",".join(parts)


CITES = Cites()


# ----------------------------------------------------------------------------------------------- run helpers
def set_font(run, name=TNR, size=9, bold=False, italic=False, sub=False, sup=False, color=None, raise_pt=None):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    run.font.size = Pt(size)
    run.font.bold = bold or None
    run.font.italic = italic or None
    if sub:
        run.font.subscript = True
    if sup:
        run.font.superscript = True
    if color is not None:
        run.font.color.rgb = color
    if raise_pt is not None:                       # the reference raises citation numerals with w:position (half-points)
        pos = OxmlElement("w:position"); pos.set(qn("w:val"), str(int(raise_pt * 2))); rpr.append(pos)


TAG = re.compile(r"(<i>|</i>|<b>|</b>|<sup>|</sup>|<sub>|</sub>|<ref>|</ref>|\[\[[a-z0-9_,]+\]\])")


def rich(par, text: str, font=TNR, size=9, bold=False, italic=False, cite_size=7, cite_raise=3):
    text = subst(text)
    st = dict(i=italic, b=bold, sup=False, sub=False, ref=False)
    for tok in TAG.split(text):
        if not tok:
            continue
        if tok in ("<i>", "</i>"):
            st["i"] = tok == "<i>" or italic
        elif tok in ("<b>", "</b>"):
            st["b"] = tok == "<b>" or bold
        elif tok in ("<sup>", "</sup>"):
            st["sup"] = tok == "<sup>"
        elif tok in ("<sub>", "</sub>"):
            st["sub"] = tok == "<sub>"
        elif tok in ("<ref>", "</ref>"):
            st["ref"] = tok == "<ref>"
        elif tok.startswith("[["):
            set_font(par.add_run(CITES.render(tok[2:-2].split(","))), font if font != CORBEL else CORBEL, cite_size, color=BLUE, raise_pt=cite_raise)
        else:
            set_font(par.add_run(tok), font, size, st["b"], st["i"], st["sub"], st["sup"], color=BLUE if st["ref"] else None)


# ----------------------------------------------------------------------------------------------- builder
class Builder:
    def __init__(self):
        self.doc = d = docx.Document(REF_DOCX)
        self.tblpr = copy.deepcopy(d.tables[3]._tbl.tblPr)
        self.authors_p = copy.deepcopy(d.paragraphs[1]._p)
        self.affil_p = copy.deepcopy(d.paragraphs[4]._p)
        body = d.element.body
        for ch in list(body):
            if ch.tag != qn("w:sectPr"):
                body.remove(ch)
        for para in d.sections[0].header.paragraphs:        # running title; the footer (authors, page number) is kept as is
            if para.runs:
                para.runs[0].text = RUNNING_TITLE
                for r in para.runs[1:]:
                    r.text = ""
        self.first_body = True
        self.ntab = self.nfig = self.nalg = self.stab = self.sfig = 0
        self.in_supp = False

    def _append(self, el):
        body = self.doc.element.body
        body.insert(list(body).index(body.find(qn("w:sectPr"))), el)

    # -- front matter (author and affiliation blocks are the reference's own paragraphs: same authors)
    def front(self):
        d = self.doc
        p = d.add_paragraph(style="SR Title")
        p.paragraph_format.space_before = Pt(74)
        set_font(p.add_run(TITLE), CORBEL, 26, True)
        self._append(self.authors_p)
        p = d.add_paragraph(style="SR Abstract")
        rich(p, ABSTRACT, CORBEL, 9, bold=True)
        p = d.add_paragraph(style="SR Keywords")
        set_font(p.add_run("Keywords"), CORBEL, 10, True)
        set_font(p.add_run("  " + KEYWORDS), TNR, 9)
        fp = self.affil_p.find(qn("w:pPr")).find(qn("w:framePr"))
        for a in ("w:y",):                                   # anchor the block to the bottom of the text area so no body line falls below it
            if fp.get(qn(a)) is not None:
                del fp.attrib[qn(a)]
        fp.set(qn("w:vAnchor"), "margin"); fp.set(qn("w:yAlign"), "bottom")
        self._append(self.affil_p)

    # -- text
    def _head(self, style, t, size):
        p = self.doc.add_paragraph(style=style)
        set_font(p.add_run(t), CORBEL, size, True)
        p.paragraph_format.keep_with_next = True
        self.first_body = True

    def h1(self, t):
        self._head("SR Heading 1", t, 11)

    def h2(self, t):
        self._head("SR Heading 2", t, 10)

    def para(self, text):
        p = self.doc.add_paragraph(style="SR Body first" if self.first_body else "Normal")
        rich(p, text)
        self.first_body = False
        return p

    def bullet(self, text):
        p = self.doc.add_paragraph(style="SR Bullet")
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Emu(198 * 635))
        p.add_run("\u2022\t")
        rich(p, text)
        return p

    def ref(self, n, text):
        p = self.doc.add_paragraph(style="SR Reference")
        rich(p, f"{n}.\t{text}", TNR, 8)

    def algorithm(self, lines, caption):
        """lines: (level, text); level 0 = BEGIN / END / step header (bold italic), level 1 = statement (italic)."""
        for level, text in lines:
            p = self.doc.add_paragraph(style="SR Algorithm")
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.left_indent = Emu((440 if level == 0 else 568 + 128 * (level - 1)) * 635)
            set_font(p.add_run(subst(text)), CAMBRIA, 8, bold=(level == 0), italic=True)
        self.nalg += 1
        p = self.doc.add_paragraph(style="SR Caption")
        set_font(p.add_run(f"Algorithm {self.nalg}:"), TNR, 9, True)
        p.add_run(" ")
        rich(p, caption, TNR, 9)
        self.first_body = False

    # -- figures and tables (caption below, as in the reference)
    def figure(self, name, caption, width=4.17, wide=False):
        p = self.doc.add_paragraph(style="SR Figure")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if wide:                                             # diagrams span the page: start in the left margin, 7.06 in wide
            p.paragraph_format.left_indent = Emu(-2180 * 635)
            width = 7.06
        p.add_run().add_picture(str(ROOT / "figures" / name), width=Inches(width))
        if self.in_supp:
            self.sfig += 1; lab = f"Supplementary Fig. S{self.sfig}."
        else:
            self.nfig += 1; lab = f"Fig. {self.nfig}."
        self.caption(lab, caption)

    def caption(self, label, text):
        p = self.doc.add_paragraph(style="SR Caption")
        set_font(p.add_run(label), TNR, 9, True)
        p.add_run("  ")
        rich(p, text, TNR, 9)
        self.first_body = False

    def table(self, header, rows, widths, caption, align=None, bold_first_col=False, top=False):
        d = self.doc
        n = len(header)
        w = [int(TEXT_W * x / sum(widths)) for x in widths]
        w[0] += TEXT_W - sum(w)
        align = align or (["left"] + ["center"] * (n - 1))
        allrows = [header] + rows
        t = d.add_table(rows=len(allrows), cols=n)
        tbl = t._tbl
        tbl.remove(tbl.tblPr)
        pr = copy.deepcopy(self.tblpr)
        for el in pr.findall(qn("w:tblLayout"))[1:]:
            pr.remove(el)
        pr.find(qn("w:tblW")).set(qn("w:w"), str(TEXT_W))
        tbl.insert(0, pr)
        for gc, wi in zip(tbl.tblGrid.findall(qn("w:gridCol")), w):
            gc.set(qn("w:w"), str(wi))
        for ri, rowvals in enumerate(allrows):
            row = t.rows[ri]
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for ci_, (cell, val) in enumerate(zip(row.cells, rowvals)):
                tcpr = cell._tc.get_or_add_tcPr()
                for el in list(tcpr):
                    tcpr.remove(el)
                tcw = OxmlElement("w:tcW"); tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(w[ci_])); tcpr.append(tcw)
                if ri == 0:
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "A8A9AC"); tcpr.append(shd)
                va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "top" if (top and ri) else "center"); tcpr.append(va)
                p = cell.paragraphs[0]
                p.style = d.styles["SR Table text"]
                p.paragraph_format.keep_with_next = True
                al = align[ci_] if ri else ("left" if (ci_ == 0 or top) else "center")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if al == "left" else WD_ALIGN_PARAGRAPH.CENTER
                rich(p, str(val), TNR, 7, bold=(ri == 0) or (bold_first_col and ci_ == 0), cite_size=5.5, cite_raise=2)
        if self.in_supp:
            self.stab += 1; lab = f"Supplementary Table S{self.stab}."
        else:
            self.ntab += 1; lab = f"Table {self.ntab}."
        self.caption(lab, caption)


LABELS: dict[str, str] = {}


def assign_labels(items) -> None:
    """Number figures, tables and algorithms in order of appearance so that {@label} references cannot drift."""
    nf = nt = na = sf = st = 0
    supp = False
    for it in items:
        k = it[0]
        if k == "supp":
            supp = True
        elif k in ("fig", "figwide"):
            if supp:
                sf += 1; LABELS[it[-1]] = f"Supplementary Fig. S{sf}"
            else:
                nf += 1; LABELS[it[-1]] = f"Fig. {nf}"
        elif k == "tab":
            if supp:
                st += 1; LABELS[it[-1]] = f"Supplementary Table S{st}"
            else:
                nt += 1; LABELS[it[-1]] = f"Table {nt}"
        elif k == "alg":
            na += 1; LABELS[it[-1]] = f"Algorithm {na}"


def xref(text: str) -> str:
    def rep(m):
        if m.group(1) not in LABELS:
            raise KeyError(f"unknown cross-reference {m.group(1)}")
        return f"<ref>{LABELS[m.group(1)]}</ref>"
    return re.sub(r"\{@([a-z0-9_:]+)\}", rep, text)


def run_items(B: Builder, items, ctx):
    for kind, *args in items:
        if kind == "h1":
            B.h1(args[0])
        elif kind == "h2":
            B.h2(args[0])
        elif kind == "p":
            B.para(xref(args[0]))
        elif kind == "b":
            B.bullet(xref(args[0]))
        elif kind == "fig":
            B.figure(args[0], xref(args[1]), args[2])
        elif kind == "figwide":
            B.figure(args[0], xref(args[1]), wide=True)
        elif kind == "alg":
            B.algorithm(args[0], xref(args[1]))
        elif kind == "tab":
            header, rows, widths, caption, *rest = args[0](ctx)
            B.table(header, rows, widths, xref(caption), *rest)
        elif kind == "supp":
            B.in_supp = True
            B.first_body = True
        else:
            raise ValueError(kind)


def main():
    assign_labels(list(BODY) + list(BACK))
    B = Builder()
    B.front()
    ctx = context()
    run_items(B, BODY, ctx)
    B.h1("References")
    for i, k in enumerate(CITES.order, 1):
        B.ref(i, REFS[k])
    unused = set(REFS) - set(CITES.order)
    assert not unused, f"uncited references: {unused}"
    nref = len(CITES.order)
    run_items(B, BACK, ctx)
    assert len(CITES.order) == nref, "a citation first appears after the reference list"
    B.doc.core_properties.title = TITLE
    B.doc.core_properties.author = "Ansh Jerath; Jagadeesan S"
    B.doc.core_properties.comments = ""
    import sys
    out = OUT.with_name(sys.argv[1]) if len(sys.argv) > 1 else OUT      # optional alternative file name (e.g. when the default is open in Word)
    B.doc.save(str(out))
    print("saved", out.name, "| tables", B.ntab, "figures", B.nfig, "algorithms", B.nalg, "| supp tables", B.stab, "supp figures", B.sfig, "| refs", nref)


if __name__ == "__main__":
    main()
