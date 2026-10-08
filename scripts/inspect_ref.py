"""Inspect formatting details of the reference manuscript (algorithm lines, review table, bullets, body runs)."""
import re
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

from src.utils import ROOT

REF = str(ROOT / "paper" / "_template.docx")   # local, git-ignored template
d = docx.Document(REF)


def show(el, n=1200):
    s = etree.tostring(el, pretty_print=True).decode()
    s = re.sub(r'xmlns:\w+="[^"]+"\s*', "", s)
    sys.stdout.buffer.write((s[:n] + "\n").encode("utf-8"))


P = d.paragraphs
for i in (70, 71, 73, 85):
    print("--- algorithm para", i)
    show(P[i]._p, 700)
print("--- bullet para 8"); show(P[8]._p, 1500)
print("--- body-first para 15 (italic lead)"); show(P[15]._p, 1100)
t = d.tables[0]
print("--- table 0 grid:", [g.get(qn("w:w")) for g in t._tbl.tblGrid.findall(qn("w:gridCol"))])
show(t.rows[1].cells[1]._tc, 900)
print("--- table gap para after table 0")
body = d.element.body
kids = list(body.iterchildren())
ix = kids.index(t._tbl)
show(kids[ix + 1], 600)
print("--- SR Caption for algorithm (94)"); show(P[94]._p, 700)
print("--- section sectPr"); show(d.sections[0]._sectPr, 1500)
